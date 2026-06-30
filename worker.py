"""
SmartGridAutomation — Playwright Scraper Worker
Logs into smartgridsoft.in, scrapes Profile + MyMeter data for each resident.
"""
import sys
import asyncio

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

import os
import json
import csv
import logging
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from playwright.async_api import async_playwright
from database import SessionLocal, Job, ResidentData

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).parent
LOGS_DIR = BASE_DIR / "logs"
LOGS_DIR.mkdir(exist_ok=True)
SCREENSHOTS_DIR = BASE_DIR / "screenshots"
SCREENSHOTS_DIR.mkdir(exist_ok=True)
DOWNLOADS_DIR = BASE_DIR / "downloads"
DOWNLOADS_DIR.mkdir(exist_ok=True)
EXPORTS_DIR = BASE_DIR / "exports"
EXPORTS_DIR.mkdir(exist_ok=True)

# Configure file + console logging
logger = logging.getLogger("smartgrid-worker")
logger.setLevel(logging.INFO)
if not logger.handlers:
    fh = logging.FileHandler(LOGS_DIR / "automation.log", encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    ch = logging.StreamHandler()
    ch.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    logger.addHandler(fh)
    logger.addHandler(ch)

# ---------------------------------------------------------------------------
# Society value map (label -> dropdown value on the real site)
# ---------------------------------------------------------------------------
SOCIETY_VALUES = {
    "Bharat City": "21",
    "Olive County": "53",
    "Gulshan Ikebana": "23",
    "Exotica Dreamville": "24",
    "Orbit Plaza": "25",
    "Cherry County": "26",
    "Shalimar City": "35",
    "Nirala Estate": "36",
    "Bharat City Phase-2": "52",
    "JYOTI SUPER VILLAGE": "38",
    "AJNARA HOMES": "39",
    "RAJNAGAR REIDENCY": "40",
    "Cleo County": "44",
    "RAJ Empire": "42",
    "Addela": "43",
    "Windsor Court": "47",
    "Ajnara Daffodil": "48",
    "Sukh Sagar Complex": "50",
    "Hawelia Valencia Homes": "51",
    "Palm Resort": "54",
    "Shalimar City RG": "55",
    "Hawelia Valenova Park": "56",
    "Green Arch": "57",
    "COCO County": "58",
    "County 107": "59",
    "IVY County": "60",
}

LOGIN_URL = "https://www.smartgridsoft.in/WebsitePages/Login.aspx"
METER_URL = "https://www.smartgridsoft.in/WebsitePages/MyMeter.aspx"
MAX_RETRIES = 3
COOLDOWN = 1.5  # seconds between accounts


# ---------------------------------------------------------------------------
# Helper: safe text extraction
# ---------------------------------------------------------------------------
async def safe_text(page, selector: str, default: str = "") -> str:
    """Safely extract text from a selector. Returns default if not found."""
    try:
        el = page.locator(selector)
        if await el.count() > 0:
            return (await el.inner_text()).strip()
    except Exception:
        pass
    return default


# ---------------------------------------------------------------------------
# Scrape Profile Page
# ---------------------------------------------------------------------------
async def scrape_profile(page) -> dict:
    """Extract resident and property info from Profile.aspx using JavaScript."""
    logger.info("  Scraping Profile page...")
    profile_data = await page.evaluate("""
    () => {
        const getTableValue = (labelText) => {
            const trs = document.querySelectorAll('tr');
            const target = labelText.toLowerCase().trim();
            for (const tr of trs) {
                const tds = tr.querySelectorAll('td');
                if (tds.length >= 3) {
                    const label = tds[1].innerText.trim().toLowerCase();
                    const cleanLabel = label.replace(/[.:]/g, '').trim();
                    if (cleanLabel === target || 
                        (target === 'email' && cleanLabel === 'email') ||
                        (target === 'alternate email' && cleanLabel === 'alternate email') ||
                        (target === 'mobile no' && cleanLabel === 'mobile no') ||
                        (target === 'alternate mobile' && cleanLabel === 'alternate mobile no') ||
                        (target === 'anniversary' && cleanLabel === 'date of annivesary') ||
                        (target === 'anniversary' && cleanLabel === 'date of anniversary')) {
                        return tds[tds.length - 1].innerText.trim();
                    }
                }
            }
            return '';
        };

        const getParagraphValue = (labelText) => {
            const ps = document.querySelectorAll('p');
            const target = labelText.toLowerCase().trim();
            for (const p of ps) {
                const text = p.innerText.trim();
                if (text.includes(':')) {
                    const parts = text.split(':');
                    const label = parts[0].trim().toLowerCase().replace(/[.:]/g, '');
                    if (label === target || 
                        (target === 'tower' && label === 'tower no') ||
                        (target === 'flat' && label === 'flat no') ||
                        (target === 'registry' && label === 'date of registry') ||
                        (target === 'possession' && label === 'date of possession') ||
                        (target === 'car 1' && label === 'car 1 no') ||
                        (target === 'car 2' && label === 'car 2 no') ||
                        (target === 'parking 1' && label === 'parking 1 no') ||
                        (target === 'parking 2' && label === 'parking 2 no')) {
                        return parts.slice(1).join(':').trim();
                    }
                }
            }
            return '';
        };

        // Try to get the name from the profile heading
        let name = '';
        const h3s = document.querySelectorAll('h3');
        for (const h3 of h3s) {
            const txt = h3.innerText.trim();
            if (txt && !txt.toLowerCase().includes('my profile') && !txt.toLowerCase().includes('my family')) {
                name = txt;
                break;
            }
        }

        return {
            name: name,
            email: getTableValue('Email'),
            alternate_email: getTableValue('Alternate Email'),
            mobile: getTableValue('Mobile No'),
            alternate_mobile: getTableValue('Alternate Mobile'),
            profession: getTableValue('Profession'),
            status: getTableValue('Status'),
            blood_group: getTableValue('Blood Group'),
            occupied_type: getTableValue('Occupied Type'),
            date_of_birth: getTableValue('Date Of Birth'),
            anniversary: getTableValue('Annivesary') || getTableValue('Anniversary'),
            society_name: getParagraphValue('Society Name'),
            tower: getParagraphValue('Tower'),
            flat_no: getParagraphValue('Flat'),
            registry_date: getParagraphValue('Registry'),
            possession_date: getParagraphValue('Possession'),
            car_1: getParagraphValue('Car 1'),
            car_2: getParagraphValue('Car 2'),
            parking_1: getParagraphValue('Parking 1'),
            parking_2: getParagraphValue('Parking 2'),
        };
    }
    """)
    return profile_data


# ---------------------------------------------------------------------------
# Scrape MyMeter Page
# ---------------------------------------------------------------------------
async def scrape_meter(page) -> dict:
    """Extract ALL meter data from MyMeter.aspx."""
    logger.info("  Scraping MyMeter page...")
    data = {}

    # --- Dashboard Cards ---
    data["balance"] = await safe_text(page, "#ContentPlaceHolder1_lblbalance")
    data["eb_units"] = await safe_text(page, "#ContentPlaceHolder1_lblebunit")
    data["eb_unit_type"] = await safe_text(page, "#ContentPlaceHolder1_lbl_EBUnitsLoad")
    data["dg_units"] = await safe_text(page, "#ContentPlaceHolder1_lbldgunit")
    data["dg_unit_type"] = await safe_text(page, "#ContentPlaceHolder1_lbl_DGUnitLoad")
    data["solar_units"] = await safe_text(page, "#ContentPlaceHolder1_lblsolarunit")
    data["wind_units"] = await safe_text(page, "#ContentPlaceHolder1_lblwindunit")
    data["meter_status"] = await safe_text(page, "#ContentPlaceHolder1_lblStatus")

    # --- Monthly Consumption (up to 6 months shown) ---
    monthly = []
    for suffix in ["one", "two", "three", "four", "five", "six"]:
        month = await safe_text(page, f"#ContentPlaceHolder1_lblmonth_{suffix}")
        amount = await safe_text(page, f"#ContentPlaceHolder1_lblmonth_{suffix}_amount")
        if month and amount:
            monthly.append({"month": month, "amount": amount})
    data["monthly_consumption"] = monthly

    # --- Applicable Rates ---
    data["applicable_rates"] = {
        "eb_rate": await safe_text(page, "#ContentPlaceHolder1_lbl_arebrate"),
        "eb_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_arebunit"),
        "dg_rate": await safe_text(page, "#ContentPlaceHolder1_lbl_ardgrate"),
        "dg_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_ardgunit"),
        "solar_rate": await safe_text(page, "#ContentPlaceHolder1_lbl_arsolarrate"),
        "solar_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_arsolarunit"),
        "wind_rate": await safe_text(page, "#ContentPlaceHolder1_lbl_arwindrate"),
        "wind_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_arwindunit"),
        "fix_charge_rate": await safe_text(page, "#ContentPlaceHolder1_lbl_arfixcrate"),
        "fix_charge_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_arfixcunit"),
    }

    # --- Source Changeover ---
    data["source_changeover"] = {
        "changeover": await safe_text(page, "#ContentPlaceHolder1_lbl_chgsrc"),
        "phase_s1": await safe_text(page, "#ContentPlaceHolder1_lbl_ebphase"),
        "phase_s2": await safe_text(page, "#ContentPlaceHolder1_lbl_dgphase"),
        "phase_s3": await safe_text(page, "#ContentPlaceHolder1_lbl_sophase"),
        "phase_s4": await safe_text(page, "#ContentPlaceHolder1_lbl_usphase"),
    }

    # --- Sanction Load ---
    data["sanction_load"] = {
        "auto_cut": await safe_text(page, "#ContentPlaceHolder1_lbl_md_cut"),
        "s1_load": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_ebsload"),
        "s1_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_ebsloadunit"),
        "s2_load": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_dgsload"),
        "s2_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_dgsloadunit"),
        "s3_load": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_sosload"),
        "s3_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_sosloadunit"),
        "s4_load": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_ussload"),
        "s4_unit": await safe_text(page, "#ContentPlaceHolder1_lbl_mod_ussloadunit"),
    }

    # --- Free Monthly Units ---
    data["free_monthly_units"] = {
        "eb": await safe_text(page, "#ContentPlaceHolder1_lbl_free_eb"),
        "dg": await safe_text(page, "#ContentPlaceHolder1_lbl_free_dg"),
        "solar": await safe_text(page, "#ContentPlaceHolder1_lbl_free_solar"),
        "wind": await safe_text(page, "#ContentPlaceHolder1_lbl_free_wind"),
    }

    # --- Current Day Deduction ---
    data["current_day_deduction"] = {
        "total": await safe_text(page, "#ContentPlaceHolder1_lbl_cdd_total"),
        "eb": await safe_text(page, "#ContentPlaceHolder1_lbl_cddeb"),
        "dg": await safe_text(page, "#ContentPlaceHolder1_lbl_cdddg"),
        "solar": await safe_text(page, "#ContentPlaceHolder1_lbl_cddsolar"),
        "wind": await safe_text(page, "#ContentPlaceHolder1_lbl_cddwind"),
        "fix_charge": await safe_text(page, "#ContentPlaceHolder1_lblcddfixc"),
    }

    # --- Previous Day Deduction ---
    data["previous_day_deduction"] = {
        "total": await safe_text(page, "#ContentPlaceHolder1_lbl_pddtotal"),
        "eb": await safe_text(page, "#ContentPlaceHolder1_lbl_pddeb"),
        "dg": await safe_text(page, "#ContentPlaceHolder1_lbl_pdddg"),
        "solar": await safe_text(page, "#ContentPlaceHolder1_lbl_pddsolar"),
        "wind": await safe_text(page, "#ContentPlaceHolder1_lbl_pddwind"),
        "fix_charge": await safe_text(page, "#ContentPlaceHolder1_lbl_pddfixc"),
    }

    # --- Last 10 Recharges (dynamic table) ---
    recharges = await page.evaluate("""
    () => {
        const rows = [];
        // Find the table that contains "Amount(INR)" header
        const tables = document.querySelectorAll('table');
        for (const table of tables) {
            const headerRow = table.querySelector('thead tr, tr:first-child');
            if (headerRow && headerRow.innerText.includes('Amount')) {
                const tbody = table.querySelector('tbody') || table;
                const trs = tbody.querySelectorAll('tr');
                for (const tr of trs) {
                    const tds = tr.querySelectorAll('td');
                    if (tds.length >= 3) {
                        rows.push({
                            amount: tds[0].innerText.trim(),
                            date: tds[1].innerText.trim(),
                            type: tds[2].innerText.trim()
                        });
                    }
                }
            }
        }
        return rows;
    }
    """)
    data["last_10_recharges"] = recharges

    # --- Meter Information ---
    data["meter_info"] = {
        "serial_number": await safe_text(page, "#ContentPlaceHolder1_lblSrN"),
        "remark": await safe_text(page, "#ContentPlaceHolder1_lblRemark"),
    }

    # --- Happy Hours ---
    happy_hours = []
    for i in range(1, 5):
        start = await safe_text(page, f"#ContentPlaceHolder1_lbl_hpsptime{i}")
        end = await safe_text(page, f"#ContentPlaceHolder1_lbl_hpeptime{i}")
        if start or end:
            happy_hours.append({"schedule": f"Schedule {i}", "start": start, "end": end})
    data["happy_hours"] = happy_hours

    # --- Happy Days ---
    happy_days = []
    for i in range(1, 32):
        val = await safe_text(page, f"#ContentPlaceHolder1_lbl_hpday{i}")
        if val:
            happy_days.append({"day": f"Happy Day {i}", "date": val})
    data["happy_days"] = happy_days

    # --- Current Month Unit Consumption (dynamic table at bottom) ---
    consumption = await page.evaluate("""
    () => {
        const rows = [];
        // This table usually has Date, EB Reading, EB Consumption, DG Reading, etc.
        const tables = document.querySelectorAll('table');
        for (const table of tables) {
            const headerRow = table.querySelector('thead tr, tr:first-child');
            if (headerRow && headerRow.innerText.includes('EB') && headerRow.innerText.includes('DG') && headerRow.innerText.includes('Date')) {
                const allRows = table.querySelectorAll('tbody tr, tr');
                for (let i = 1; i < allRows.length; i++) {  // skip header
                    const tds = allRows[i].querySelectorAll('td');
                    if (tds.length >= 6) {
                        rows.push({
                            date: tds[0].innerText.trim(),
                            eb_reading: tds[1].innerText.trim(),
                            eb_consumption: tds[2].innerText.trim(),
                            dg_reading: tds[3].innerText.trim(),
                            dg_consumption: tds[4].innerText.trim(),
                            balance: tds[5].innerText.trim(),
                            time: tds.length > 6 ? tds[6].innerText.trim() : ''
                        });
                    }
                }
            }
        }
        return rows;
    }
    """)
    data["current_month_consumption"] = consumption

    # --- Source Running & Last Sync Info ---
    data["last_sync"] = await page.evaluate("""
    () => {
        // Look for the red text showing last updated date/time
        const spans = document.querySelectorAll('span');
        for (const span of spans) {
            const text = span.innerText.trim();
            if (text.includes('Last Updated') || text.includes('(')) {
                // e.g. "(30-06-2026 00:10:24)"
                const match = text.match(/\\((\\d{2}-\\d{2}-\\d{4}\\s+\\d{2}:\\d{2}:\\d{2})\\)/);
                if (match) return match[1];
            }
        }
        return '';
    }
    """)

    # --- Meter Date & Time ---
    data["meter_date_time"] = {
        "meter_date": await safe_text(page, "#ContentPlaceHolder1_lbl_meterdate",
                                       await safe_text(page, "span[id*='meterdate']")),
        "meter_time": await safe_text(page, "#ContentPlaceHolder1_lbl_metertime",
                                       await safe_text(page, "span[id*='metertime']")),
    }

    return data


# ---------------------------------------------------------------------------
# Process one resident account
# ---------------------------------------------------------------------------
async def process_account(
    page, context, society_name: str, username: str, password: str, job_id: str
) -> dict:
    """
    Login → Scrape Profile → Scrape MyMeter → Return all data.
    Retries up to MAX_RETRIES on transient failures.
    """
    society_value = SOCIETY_VALUES.get(society_name)
    if not society_value:
        return {
            "status": "Error",
            "error": f"Unknown society: {society_name}",
            "data": None,
        }

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            # Reset state
            await context.clear_cookies()
            if not hasattr(page, "dialog_messages"):
                page.dialog_messages = []
                async def on_dialog(dialog):
                    logger.info(f"    Dialog: [{dialog.type}] '{dialog.message}'")
                    page.dialog_messages.append(dialog.message)
                    await dialog.dismiss()
                page.on("dialog", on_dialog)
            else:
                page.dialog_messages.clear()

            # --- Step 1: Navigate to login ---
            logger.info(f"  [{attempt}/{MAX_RETRIES}] Navigating to login for {username}...")
            await page.goto(LOGIN_URL, timeout=60000, wait_until="domcontentloaded")
            await page.wait_for_load_state("networkidle")

            # --- Step 2: Select society ---
            logger.info(f"  Selecting society: {society_name} (value={society_value})")
            await page.select_option("select#ddlsocietyname", value=society_value)
            await page.wait_for_load_state("networkidle")

            # --- Step 3: Enter credentials ---
            logger.info(f"  Entering username: {username}")
            await page.fill("input#txtusername", "")
            await page.fill("input#txtusername", str(username))
            await page.fill("input#txtpassword", "")
            await page.fill("input#txtpassword", str(password))

            # --- Step 4: Click Sign In ---
            logger.info("  Clicking Sign In...")
            await page.click("input#btnsignin")

            # Wait for redirect or alert
            for _ in range(30):  # 15 second timeout
                if page.dialog_messages:
                    return {
                        "status": "Auth Failed",
                        "error": page.dialog_messages[0],
                        "data": None,
                    }
                if "Profile.aspx" in page.url or "MyMeter.aspx" in page.url:
                    break
                await asyncio.sleep(0.5)
            else:
                # Timeout — neither redirect nor dialog
                if page.dialog_messages:
                    return {
                        "status": "Auth Failed",
                        "error": page.dialog_messages[0],
                        "data": None,
                    }
                # Take screenshot
                ss_path = SCREENSHOTS_DIR / f"{job_id}_{username}_{password}_timeout.png"
                await page.screenshot(path=str(ss_path))
                if attempt < MAX_RETRIES:
                    logger.warning(f"  Timeout for {username}, retrying...")
                    continue
                return {
                    "status": "Error",
                    "error": "Login timed out — no redirect or error detected",
                    "data": None,
                }

            logger.info(f"  Login successful! Redirected to: {page.url}")
            await page.wait_for_load_state("networkidle")
            await page.wait_for_selector(".profile-data, h3", timeout=15000)
            await asyncio.sleep(1)

            # --- Step 5: Scrape Profile ---
            profile_data = await scrape_profile(page)

            # --- Step 6: Navigate to MyMeter and scrape ---
            await page.goto(METER_URL, timeout=60000, wait_until="domcontentloaded")
            await page.wait_for_load_state("networkidle")
            await asyncio.sleep(2)  # Let dynamic data load

            meter_data = await scrape_meter(page)

            # Combine all data
            full_data = {
                "profile": profile_data,
                "meter": meter_data,
                "scraped_at": datetime.now(timezone.utc).isoformat(),
            }

            return {
                "status": "Success",
                "error": None,
                "data": full_data,
            }

        except Exception as e:
            logger.error(f"  Error on attempt {attempt} for {username}: {e}")
            try:
                ss_path = SCREENSHOTS_DIR / f"{job_id}_{username}_{password}_error_{attempt}.png"
                await page.screenshot(path=str(ss_path))
            except Exception:
                pass

            if attempt < MAX_RETRIES:
                retry_delay = attempt * 5
                logger.info(f"  Retrying in {retry_delay} seconds...")
                await asyncio.sleep(retry_delay)
                continue

            return {
                "status": "Error",
                "error": str(e),
                "data": None,
            }


# ---------------------------------------------------------------------------
# Export functions
# ---------------------------------------------------------------------------
def generate_exports(job_id: str, society_name: str, db):
    """Generate Excel, CSV, JSON, and failed_logins.csv exports."""
    records = db.query(ResidentData).filter(ResidentData.job_id == job_id).all()
    job_exports_dir = EXPORTS_DIR / job_id[:8]
    job_exports_dir.mkdir(exist_ok=True)

    all_data = []
    failed_rows = []

    for r in records:
        if r.status == "Success" and r.data_json:
            parsed = json.loads(r.data_json)
            parsed["_username"] = r.username
            parsed["_password"] = r.password
            parsed["_status"] = r.status
            all_data.append(parsed)
        elif r.status != "Pending":
            failed_rows.append({
                "username": r.username,
                "password": r.password,
                "status": r.status,
                "error": r.error_message or "",
                "society": society_name,
                "timestamp": r.scraped_at.isoformat() if r.scraped_at else "",
            })

    # ---- JSON Export ----
    json_path = job_exports_dir / "Residents.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(all_data, f, indent=2, ensure_ascii=False)

    # ---- CSV Export (flat view of key fields) ----
    csv_path = job_exports_dir / "Residents.csv"
    csv_headers = [
        "Username", "Password", "Name", "Email", "Alternate Email", 
        "Mobile", "Alternate Mobile", "Profession", "Status", "Blood Group", 
        "Occupied Type", "Date Of Birth", "Anniversary", "Society Name", 
        "Tower", "Flat", "Registry Date", "Possession Date", 
        "Car 1", "Car 2", "Parking 1", "Parking 2",
        "Balance", "EB_Units", "DG_Units", "Solar_Units", "Wind_Units", "Meter_Status",
        "EB_Rate", "DG_Rate", "Fix_Charge_Per_Day",
        "Current_Day_Total", "Previous_Day_Total",
        "Meter_Serial", "Meter_Remark",
    ]
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=csv_headers)
        writer.writeheader()
        for d in all_data:
            profile = d.get("profile", {})
            meter = d.get("meter", {})
            rates = meter.get("applicable_rates", {})
            cdd = meter.get("current_day_deduction", {})
            pdd = meter.get("previous_day_deduction", {})
            mi = meter.get("meter_info", {})
            writer.writerow({
                "Username": d["_username"],
                "Password": d["_password"],
                "Name": profile.get("name", ""),
                "Email": profile.get("email", ""),
                "Alternate Email": profile.get("alternate_email", ""),
                "Mobile": profile.get("mobile", ""),
                "Alternate Mobile": profile.get("alternate_mobile", ""),
                "Profession": profile.get("profession", ""),
                "Status": profile.get("status", ""),
                "Blood Group": profile.get("blood_group", ""),
                "Occupied Type": profile.get("occupied_type", ""),
                "Date Of Birth": profile.get("date_of_birth", ""),
                "Anniversary": profile.get("anniversary", ""),
                "Society Name": profile.get("society_name", ""),
                "Tower": profile.get("tower", ""),
                "Flat": profile.get("flat_no", ""),
                "Registry Date": profile.get("registry_date", ""),
                "Possession Date": profile.get("possession_date", ""),
                "Car 1": profile.get("car_1", ""),
                "Car 2": profile.get("car_2", ""),
                "Parking 1": profile.get("parking_1", ""),
                "Parking 2": profile.get("parking_2", ""),
                "Balance": meter.get("balance", ""),
                "EB_Units": meter.get("eb_units", ""),
                "DG_Units": meter.get("dg_units", ""),
                "Solar_Units": meter.get("solar_units", ""),
                "Wind_Units": meter.get("wind_units", ""),
                "Meter_Status": meter.get("meter_status", ""),
                "EB_Rate": rates.get("eb_rate", ""),
                "DG_Rate": rates.get("dg_rate", ""),
                "Fix_Charge_Per_Day": rates.get("fix_charge_rate", ""),
                "Current_Day_Total": cdd.get("total", ""),
                "Previous_Day_Total": pdd.get("total", ""),
                "Meter_Serial": mi.get("serial_number", ""),
                "Meter_Remark": mi.get("remark", ""),
            })

    # ---- Excel Export ----
    xlsx_path = job_exports_dir / "Residents.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(csv_headers)
    for d in all_data:
        profile = d.get("profile", {})
        meter = d.get("meter", {})
        rates = meter.get("applicable_rates", {})
        cdd = meter.get("current_day_deduction", {})
        pdd = meter.get("previous_day_deduction", {})
        mi = meter.get("meter_info", {})
        ws.append([
            d["_username"], d["_password"],
            profile.get("name", ""), profile.get("email", ""),
            profile.get("alternate_email", ""), profile.get("mobile", ""),
            profile.get("alternate_mobile", ""), profile.get("profession", ""),
            profile.get("status", ""), profile.get("blood_group", ""),
            profile.get("occupied_type", ""), profile.get("date_of_birth", ""),
            profile.get("anniversary", ""), profile.get("society_name", ""),
            profile.get("tower", ""), profile.get("flat_no", ""),
            profile.get("registry_date", ""), profile.get("possession_date", ""),
            profile.get("car_1", ""), profile.get("car_2", ""),
            profile.get("parking_1", ""), profile.get("parking_2", ""),
            meter.get("balance", ""), meter.get("eb_units", ""),
            meter.get("dg_units", ""), meter.get("solar_units", ""),
            meter.get("wind_units", ""), meter.get("meter_status", ""),
            rates.get("eb_rate", ""), rates.get("dg_rate", ""),
            rates.get("fix_charge_rate", ""),
            cdd.get("total", ""), pdd.get("total", ""),
            mi.get("serial_number", ""), mi.get("remark", ""),
        ])

    # Sheet 2: Monthly Consumption
    ws2 = wb.create_sheet("Monthly Consumption")
    ws2.append(["Username", "Password", "Month", "Amount"])
    for d in all_data:
        for mc in d.get("meter", {}).get("monthly_consumption", []):
            ws2.append([d["_username"], d["_password"], mc["month"], mc["amount"]])

    # Sheet 3: Recharge History
    ws3 = wb.create_sheet("Recharge History")
    ws3.append(["Username", "Password", "Amount", "Date", "Type"])
    for d in all_data:
        for rc in d.get("meter", {}).get("last_10_recharges", []):
            ws3.append([d["_username"], d["_password"], rc["amount"], rc["date"], rc["type"]])

    # Sheet 4: Current Month Consumption
    ws4 = wb.create_sheet("Current Month Consumption")
    ws4.append(["Username", "Password", "Date", "EB Reading", "EB Consumption",
                "DG Reading", "DG Consumption", "Balance", "Time"])
    for d in all_data:
        for cc in d.get("meter", {}).get("current_month_consumption", []):
            ws4.append([
                d["_username"], d["_password"],
                cc.get("date", ""), cc.get("eb_reading", ""), cc.get("eb_consumption", ""),
                cc.get("dg_reading", ""), cc.get("dg_consumption", ""),
                cc.get("balance", ""), cc.get("time", ""),
            ])

    wb.save(xlsx_path)

    # ---- Failed Logins CSV ----
    failed_csv_path = job_exports_dir / "failed_logins.csv"
    with open(failed_csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["username", "password", "status", "error", "society", "timestamp"])
        writer.writeheader()
        for row in failed_rows:
            writer.writerow(row)

    return {
        "excel": str(xlsx_path),
        "csv": str(csv_path),
        "json": str(json_path),
        "failed_csv": str(failed_csv_path),
    }


# ---------------------------------------------------------------------------
# Main automation job runner
# ---------------------------------------------------------------------------
PAUSE_WAIT_SECS = 60  # How long to wait before retrying after consecutive failures

async def run_automation_job(job_id: str):
    """
    Main entry point: process all accounts for a job.
    Supports RESUME: only processes records still in 'Pending' or 'Error' status.
    Auto-PAUSEs on persistent network failures instead of hard-aborting.
    Always generates partial exports when there's at least 1 successful record.
    """
    db = SessionLocal()
    job = db.query(Job).filter(Job.id == job_id).first()
    if not job:
        logger.error(f"Job {job_id} not found!")
        db.close()
        return

    job.status = "PROCESSING"
    db.commit()

    # --- Resume-aware record selection ---
    # Only process records that haven't been successfully scraped yet
    all_records = db.query(ResidentData).filter(ResidentData.job_id == job_id).all()
    total = len(all_records)

    # Initialize counters from existing DB state (crucial for resume)
    already_success = sum(1 for r in all_records if r.status == "Success")
    already_failed = sum(1 for r in all_records if r.status not in ("Pending", "Success"))
    already_processed = already_success + already_failed

    # Filter to only pending records (those not yet scraped or previously errored)
    pending_records = [r for r in all_records if r.status in ("Pending", "Error")]
    pending_count = len(pending_records)

    is_resume = already_processed > 0
    if is_resume:
        logger.info(f"{'='*60}")
        logger.info(f"RESUMING job {job_id} — Society: {job.society_name}")
        logger.info(f"  Already scraped: {already_success} success, {already_failed} failed")
        logger.info(f"  Remaining: {pending_count} accounts to process")
        logger.info(f"{'='*60}")
    else:
        logger.info(f"{'='*60}")
        logger.info(f"Starting job {job_id} — {total} accounts — Society: {job.society_name}")
        logger.info(f"{'='*60}")

    if pending_count == 0:
        logger.info("No pending records to process. Generating exports...")
        job.processed_records = total
        job.successful = already_success
        job.failed = already_failed
        try:
            exports = generate_exports(job_id, job.society_name, db)
            job.export_excel = exports["excel"]
            job.export_csv = exports["csv"]
            job.export_json = exports["json"]
            job.export_failed_csv = exports["failed_csv"]
            job.status = "COMPLETED"
            job.completed_at = datetime.now(timezone.utc)
        except Exception as e:
            logger.exception(f"Error generating exports: {e}")
            job.status = "FAILED"
        db.commit()
        db.close()
        return

    success_count = already_success
    fail_count = already_failed
    processed_count = already_processed

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-setuid-sandbox"]
        )
        context = await browser.new_context(
            viewport={"width": 1280, "height": 900},
            ignore_https_errors=True,
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        )
        page = await context.new_page()

        consecutive_timeouts = 0
        paused = False
        cancelled = False

        for i, record in enumerate(pending_records):
            # Check for cancellation
            db.refresh(job)
            if job.status == "CANCELLED":
                logger.info("Job was cancelled by user.")
                cancelled = True
                break

            processed_count += 1
            logger.info(f"\n--- Account {processed_count}/{total}: Username={record.username} Password={record.password} ---")

            result = await process_account(
                page=page,
                context=context,
                society_name=job.society_name,
                username=record.username,
                password=str(record.password),
                job_id=job_id,
            )

            # Update record
            record.status = result["status"]
            record.error_message = result.get("error")
            record.data_json = json.dumps(result["data"]) if result["data"] else None
            record.scraped_at = datetime.now(timezone.utc)

            if result["status"] == "Success":
                success_count += 1
                consecutive_timeouts = 0
            else:
                fail_count += 1
                # Track consecutive connection failures
                err_msg = result.get("error") or ""
                is_connection_error = (
                    "ERR_CONNECTION_TIMED_OUT" in err_msg or
                    "timed out" in err_msg.lower() or
                    "ERR_CONNECTION_REFUSED" in err_msg or
                    "ERR_NAME_NOT_RESOLVED" in err_msg or
                    "net::" in err_msg
                )
                if is_connection_error:
                    consecutive_timeouts += 1
                else:
                    consecutive_timeouts = 0  # Reset for non-network errors (auth failures etc.)

            job.processed_records = processed_count
            job.successful = success_count
            job.failed = fail_count
            db.commit()

            logger.info(f"  Result: {result['status']}" +
                        (f" — {result['error']}" if result.get('error') else ""))

            # --- Smart auto-pause with recovery ---
            if consecutive_timeouts >= 3:
                logger.warning(f"3 consecutive connection failures detected. Pausing for {PAUSE_WAIT_SECS}s to let the website recover...")
                await asyncio.sleep(PAUSE_WAIT_SECS)

                # Retry one more time to see if the site is back
                logger.info("Attempting recovery check...")
                try:
                    recovery_page = await context.new_page()
                    await recovery_page.goto(LOGIN_URL, timeout=30000, wait_until="domcontentloaded")
                    await recovery_page.close()
                    logger.info("Recovery successful! Website is back. Continuing...")
                    consecutive_timeouts = 0  # Reset and continue
                except Exception as recovery_err:
                    logger.error(f"Recovery failed: {recovery_err}")
                    logger.error("PAUSING job — The target website is persistently unreachable.")
                    logger.error("You can RESUME this job later from the portal when the website is back.")
                    paused = True
                    try:
                        await recovery_page.close()
                    except Exception:
                        pass
                    break

            # Cooldown between accounts
            if i < pending_count - 1:
                await asyncio.sleep(COOLDOWN)

        await browser.close()

    # --- Finalize: Always try to generate exports if we have data ---
    db.refresh(job)

    # Determine final status
    if cancelled or job.status == "CANCELLED":
        final_status = "CANCELLED"
    elif paused:
        final_status = "PAUSED"
    elif processed_count >= total:
        final_status = "COMPLETED"
    else:
        final_status = job.status  # Keep whatever it was

    # Always generate exports if there's at least 1 successful record
    if success_count > 0:
        try:
            exports = generate_exports(job_id, job.society_name, db)
            job.export_excel = exports["excel"]
            job.export_csv = exports["csv"]
            job.export_json = exports["json"]
            job.export_failed_csv = exports["failed_csv"]
            logger.info(f"Exports generated with {success_count} records → {EXPORTS_DIR / job_id[:8]}")
        except Exception as e:
            logger.exception(f"Error generating exports: {e}")

    if final_status == "COMPLETED":
        job.status = "COMPLETED"
        job.completed_at = datetime.now(timezone.utc)
        logger.info(f"\n{'='*60}")
        logger.info(f"Job COMPLETED — Success: {success_count}, Failed: {fail_count}")
        logger.info(f"{'='*60}")
    elif final_status == "PAUSED":
        job.status = "PAUSED"
        remaining = total - processed_count
        logger.info(f"\n{'='*60}")
        logger.info(f"Job PAUSED — Success: {success_count}, Failed: {fail_count}, Remaining: {remaining}")
        logger.info(f"Resume this job from the portal when the website is back online.")
        logger.info(f"{'='*60}")
    elif final_status == "CANCELLED":
        # Status already set by cancel API; just log it
        remaining = total - processed_count
        logger.info(f"\n{'='*60}")
        logger.info(f"Job CANCELLED — Success: {success_count}, Failed: {fail_count}, Remaining: {remaining}")
        if success_count > 0:
            logger.info(f"Partial exports available with {success_count} records.")
        logger.info(f"You can RESUME this job from the portal to process the remaining {remaining} accounts.")
        logger.info(f"{'='*60}")

    db.commit()
    db.close()


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python worker.py <job_id>")
        sys.exit(1)
    
    job_uuid = sys.argv[1]
    asyncio.run(run_automation_job(job_uuid))

